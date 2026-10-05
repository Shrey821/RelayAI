import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Colors
    BG_DARK = RGBColor(15, 20, 28)         # #0F141C
    CARD_BG = RGBColor(26, 34, 50)         # #1A2232
    CARD_BORDER = RGBColor(42, 54, 79)     # #2A364F
    ACCENT_CYAN = RGBColor(14, 165, 233)   # #0EA5E9
    ACCENT_EMERALD = RGBColor(16, 185, 129)# #10B981
    ACCENT_AMBER = RGBColor(245, 158, 11)  # #F59E0B
    ACCENT_ROSE = RGBColor(239, 68, 68)    # #EF4444
    TEXT_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8
    TEXT_DIM = RGBColor(100, 116, 139)     # #64748B
    
    # Heat map colors (accessible & readable with white text)
    HM_GREEN_STRONG = RGBColor(16, 130, 85)   # Strong / High / Optimal
    HM_GREEN_MED = RGBColor(20, 95, 65)       # Moderate-High
    HM_AMBER = RGBColor(165, 105, 12)         # Moderate / Medium
    HM_ROSE_MED = RGBColor(150, 50, 50)       # Low / Risky
    HM_ROSE_STRONG = RGBColor(180, 30, 30)    # Minimal / High Risk
    HM_SLATE = RGBColor(38, 48, 68)           # Neutral / Base
    
    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category, slide_num):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.3))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        
        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10.5), Inches(0.6))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = title
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT
        
        # Slide number
        num_box = slide.shapes.add_textbox(Inches(11.8), Inches(0.4), Inches(0.8), Inches(0.3))
        tf3 = num_box.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = f"{slide_num:02d} / 10"
        p3.alignment = PP_ALIGN.RIGHT
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_DIM
        
        # Divider line
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

    def add_card(slide, left, top, width, height, title=None, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.35))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Outer accent border frame
    frame = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD_BG
    frame.line.color.rgb = CARD_BORDER
    frame.line.width = Pt(1.5)
    
    # Glow bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.15), Inches(5.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_CYAN
    bar.line.fill.background()
    
    # Title text box
    tb = s1.shapes.add_textbox(Inches(1.4), Inches(1.5), Inches(10.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "RELAY AI  •  UNIVERSAL MEMORY SCHEMA (UMS)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(16)
    
    p2 = tf.add_paragraph()
    p2.text = "Comprehensive Business Model & SaaS Strategy"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_after = Pt(12)
    
    p3 = tf.add_paragraph()
    p3.text = "A Grounded, Subscription-First Monetization Framework & Heat Map Analysis for Cross-LLM Developer Context Portability"
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(36)
    
    p4 = tf.add_paragraph()
    p4.text = "PRICING ENGINE: Recurring Subscription (Pro, Team, Enterprise)  |  FOUNDATION: Zero-Dependency Open IR Standard\nTARGET SEGMENT: Multi-Model AI Developers & Engineering Teams  |  METHODOLOGY: No Hype, Unit-Economics-Driven"
    p4.font.size = Pt(12)
    p4.font.color.rgb = TEXT_DIM
    
    prs.save("test_run.pptx")
    print("Base generated successfully")

create_presentation()
